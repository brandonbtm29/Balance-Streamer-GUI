"""Smoke tests: the window builds, flow rates come out right, exports open in Excel."""

import datetime
import importlib.metadata

import openpyxl
import pytest

import balance_streamer


def test_version_matches_installed_package():
    try:
        installed = importlib.metadata.version("balance-streamer")
    except importlib.metadata.PackageNotFoundError:
        pytest.skip("package not installed (run: pip install -e .[dev])")
    assert installed == balance_streamer.__version__


def test_main_window_builds(window):
    assert window.tabs.count() >= 2  # Settings tab + one balance tab
    assert window.tab_objects


def _record_linear_run(tab, grams_per_min=2.0, seconds=120):
    """Fill a tab as if a balance had logged a steady flow once per second."""
    start = datetime.datetime(2026, 9, 27, 12, 56, 21)
    tab.session_start_time = start
    for t in range(seconds):
        tab.times_sec.append(float(t))
        tab.times_min.append(t / 60.0)
        tab.weights.append(10.0 + grams_per_min * t / 60.0)
        tab.timestamps.append((start + datetime.timedelta(seconds=t)).strftime("%Y-%m-%d %H:%M:%S.%f")[:-3])


@pytest.mark.parametrize("filter_type", ["Mean", "Median", "EMA", "Butterworth"])
def test_steady_flow_gives_the_right_rate(tab, filter_type):
    _record_linear_run(tab, grams_per_min=2.0)
    tab.combo_filter.setCurrentText(filter_type)
    tab.calculate_flows()
    assert tab.flow_rates[-1] == pytest.approx(2.0, rel=1e-6)


def test_filename_tokens(tab):
    tab.session_start_time = datetime.datetime(2026, 9, 27, 12, 56, 21)
    name = tab.get_resolved_string("<Start Date> <Start Time> Flow rate C1<Unknown> a/b:c")
    # Unknown tokens are dropped and characters Windows can't use become "_".
    assert name == "2026-09-27 125621 Flow rate C1 a_b_c"


def test_excel_export(tab, tmp_path):
    _record_linear_run(tab, seconds=30)
    tab.calculate_flows()
    tab.chk_use_global_template.setChecked(False)
    tab.ent_save_template.setText("<Start Date> <Start Time> test run")
    tab.ent_save_folder.setPlainText(str(tmp_path))

    tab.save_excel(auto=True)

    path = tmp_path / "2026-09-27 125621 test run.xlsx"
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Documentation", "Data", "Chart"]
    rows = list(wb["Data"].iter_rows(values_only=True))
    assert rows[0] == ("Timestamp", "Duration (min)", "Mass (g)", "Flow Rate (g/min)")
    assert len(rows) == 31
    assert rows[1][2] == pytest.approx(10.0)
