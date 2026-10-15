"""Saving a run: file-name templates, Excel workbooks and chart images."""

import datetime
import os

import openpyxl
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font
from PyQt6.QtWidgets import (
    QFileDialog,
    QMessageBox,
)


class ExportMixin:
    """Part of BalanceTab: saving a run: file-name templates, Excel workbooks and chart images."""

    def get_resolved_string(self, template_str):
        if not template_str:
            return ""
        import datetime
        import re

        s = template_str
        s = s.replace("<Tab Name>", self.tab_name)

        now = datetime.datetime.now()
        start_dt = getattr(self, "session_start_time", None) or now

        s = s.replace("<Start Date>", start_dt.strftime("%Y-%m-%d"))
        s = s.replace("<Start Time>", start_dt.strftime("%H%M%S"))
        s = s.replace("<Date>", now.strftime("%Y-%m-%d"))
        s = s.replace("<Time>", now.strftime("%H%M%S"))
        global_tokens = self.app.config.get("global_tokens", []) if self.app else []
        for t in global_tokens:
            name = t.get("name", "").strip()
            if name:
                if hasattr(self, "token_inputs") and name in self.token_inputs:
                    val = self.token_inputs[name].text()
                else:
                    val = t.get("value", "")
                s = s.replace(f"<{name}>", val)
        if hasattr(self, "token_inputs"):
            for name, widget in self.token_inputs.items():
                val = widget.text() if hasattr(widget, "text") else str(widget)
                s = s.replace(f"<{name}>", val)
        s = re.sub(r"<[^>]+>", "", s)
        s = re.sub(r'[<>:"/\\|?*]', "_", s)
        return s

    def save_excel(self, auto=False, quick_save=False):
        if not self.times_sec:
            return

        if self.chk_use_global_template.isChecked():
            template = (
                self.app.config.get("global_filename_template", "<Tab Name>_<Date>_<Time>")
                if self.app
                else "<Tab Name>_<Date>_<Time>"
            )
        else:
            template = self.ent_save_template.text()
            if not template.strip():
                template = "<Tab Name>_<Date>_<Time>"

        filename = self.get_resolved_string(template)
        if not filename.endswith(".xlsx"):
            filename += ".xlsx"

        save_dir = self.ent_save_folder.toPlainText().replace("\n", "").strip().strip()
        if not save_dir:
            save_dir = "Data"

        os.makedirs(save_dir, exist_ok=True)
        default_path = os.path.join(save_dir, filename)

        if quick_save:
            if hasattr(self, "last_saved_filepath") and self.last_saved_filepath:
                filepath = self.last_saved_filepath
            else:
                filepath = default_path
        elif auto:
            filepath = default_path
        else:
            filepath, _ = QFileDialog.getSaveFileName(
                self, "Save Excel Data", default_path, "Excel files (*.xlsx)"
            )
            if not filepath:
                return

        try:
            wb = openpyxl.Workbook()
            doc_ws = wb.active
            doc_ws.title = "Documentation"

            doc_ws["A1"] = f"{self.tab_name} - Experiment Documentation"
            doc_ws["A1"].font = Font(size=14, bold=True)
            doc_ws["A3"] = "Export Date:"
            doc_ws["B3"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            doc_ws["A4"] = "Balance Brand:"
            doc_ws["B4"] = self.combo_brand.currentText()
            doc_ws["A5"] = "COM Port:"
            doc_ws["B5"] = self.combo_com.currentText()

            def parse_auto_type(raw_str):
                if not raw_str or not isinstance(raw_str, str):
                    return raw_str
                s = raw_str.strip()
                try:
                    return int(s)
                except ValueError:
                    try:
                        return float(s)
                    except ValueError:
                        return s

            current_row = 7
            if hasattr(self, "token_inputs") and self.token_inputs:
                for name, widget in self.token_inputs.items():
                    n = name.strip()
                    if n:
                        val = widget.text() if hasattr(widget, "text") else str(widget)
                        doc_ws[f"A{current_row}"] = f"{n}:"
                        doc_ws[f"B{current_row}"] = parse_auto_type(val)
                        current_row += 1
            else:
                global_tokens = self.app.config.get("global_tokens", []) if self.app else []
                for t in global_tokens:
                    n = t.get("name", "").strip()
                    if n:
                        doc_ws[f"A{current_row}"] = f"{n}:"
                        doc_ws[f"B{current_row}"] = parse_auto_type(t.get("value", ""))
                        current_row += 1

            current_row += 1
            doc_ws[f"A{current_row}"] = "Experiment Notes:"
            doc_ws[f"A{current_row}"].font = Font(bold=True)
            doc_ws[f"A{current_row + 1}"] = self.get_resolved_string(self.txt_notes.toPlainText())
            doc_ws.column_dimensions["A"].width = 25
            doc_ws.column_dimensions["B"].width = 35

            data_ws = wb.create_sheet("Data")
            data_ws.append(["Timestamp", "Duration (min)", "Mass (g)", "Flow Rate (g/min)"])
            for cell in data_ws[1]:
                cell.font = Font(bold=True)

            for i in range(len(self.times_sec)):
                ts = str(self.timestamps[i])
                dur = float(self.times_min[i])
                mass = (
                    float(self.weights[i])
                    if (i < len(self.weights) and self.weights[i] is not None)
                    else None
                )
                flw = (
                    float(self.flow_rates[i])
                    if (i < len(self.flow_rates) and self.flow_rates[i] is not None)
                    else None
                )
                data_ws.append([ts, dur, mass, flw])

                row_idx = i + 2
                data_ws[f"B{row_idx}"].number_format = "0.00000"
                data_ws[f"C{row_idx}"].number_format = "0.000"
                if flw is not None:
                    data_ws[f"D{row_idx}"].number_format = "0.0000"

            chart_ws = wb.create_sheet("Chart")
            chart = ScatterChart()
            chart.title = f"{self.tab_name} - Weight vs. Time"
            max_row = len(self.times_sec) + 1
            if max_row > 1:
                xvalues = Reference(data_ws, min_col=2, min_row=2, max_row=max_row)
                yvalues = Reference(data_ws, min_col=3, min_row=2, max_row=max_row)
                series = Series(yvalues, xvalues, title="Recorded Weight")
                chart.series.append(series)
            chart_ws.add_chart(chart, "A1")

            wb.save(filepath)
            self.last_saved_filepath = filepath
            self.unsaved_changes = False
            if self.app:
                self.app.set_unsaved_state(self.tab_name, False)
            if auto or quick_save:
                msg = "Auto-Saved" if auto else "Quick Saved"
                self.lbl_status.setText(f"Status: {msg} at {datetime.datetime.now().strftime('%H:%M:%S')}")
                self.lbl_status.setStyleSheet("color: #27ae60; font-weight: bold;")
            else:
                QMessageBox.information(self, "Success", "Saved to Excel.")
        except Exception as e:
            if auto or quick_save:
                self.lbl_status.setText(f"Status: Save Failed ({str(e)})")
                self.lbl_status.setStyleSheet("color: #e74c3c; font-weight: bold;")
            else:
                QMessageBox.critical(self, "Error", str(e))

    def save_graph(self):
        dt_str = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
        filepath, _ = QFileDialog.getSaveFileName(
            self, "Save Graph", f"{self.tab_name.replace(' ', '_')}_graph_{dt_str}.png", "PNG Image (*.png)"
        )
        if filepath:
            try:
                self.fig.savefig(filepath)
                QMessageBox.information(self, "Success", "Graph saved.")
            except Exception as e:
                QMessageBox.critical(self, "Error", str(e))
